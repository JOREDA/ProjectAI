"""Packs the CSVs into one Excel workbook with a start page and data dictionary."""
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
D = os.path.join(HERE, "data")
OUT = os.path.join(HERE, "HR_Feedback_Synthetic_Dataset.xlsx")

ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
BODY = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
TITLE = Font(name=ARIAL, size=14, bold=True)
WARN_FILL = PatternFill("solid", fgColor="FDECEA")

sheets = [
    ("feedback_comments", "feedback_comments.csv"),
    ("team_monthly_panel", "team_monthly_panel.csv"),
    ("teams", "teams.csv"),
    ("employees_hris", "employees_hris.csv"),
    ("listening_sessions", "listening_sessions.csv"),
    ("manager_yearend_review", "manager_yearend_review.csv"),
    ("hr_review_disputes", "hr_review_disputes.csv"),
    ("identity_vault_SEALED", "identity_vault_SEALED.csv"),
    ("policy_events", "policy_events.csv"),
    ("answer_key_hidden", "answer_key_hidden_drivers.csv"),
]
frames = {name: pd.read_csv(os.path.join(D, f), keep_default_na=False, na_values=[""]) for name, f in sheets}

# ---------------------------------------------------------------- dictionary
DICT = {
 "feedback_comments": {
  "comment_id": "Unique comment ID",
  "employee_id_hidden": "Writer's employee ID. NEVER shown in the app - only used to build the sealed vault and to test re-identification",
  "team_id": "Team the writer belonged to (used only for team-level aggregation)",
  "month": "Month of submission (YYYY-MM)",
  "submission_window": "Date shown as a half-month window, never an exact date (idea: date anonymous by keeping it as a window)",
  "anonymity_mode": "Writer's choice: Sealed token (default, identity locked in vault) / Fully anonymous (no token, receipt code only) / Identified (writer chose to be named)",
  "raw_text": "What the employee typed - may contain names, project codes, office, employee ID, dates",
  "gold_layer1_text": "Correct output of privacy layer 1: direct identifiers removed (names, IDs, project codes, client, office, dates)",
  "gold_final_text": "Correct output of privacy layer 2: indirect identifiers generalised (unique roles, situations, signature phrases). This is what HR should see",
  "contains_direct_pii": "TRUE if raw_text has a direct identifier",
  "pii_types": "Which identifiers are in raw_text",
  "reid_risk_after_layer1": "Low / Medium / High - can the writer still be guessed after names are removed?",
  "reid_reason": "Why the risk is Medium/High (unique role, situation, or writing-style signature)",
  "language": "English or Hinglish",
  "gold_theme": "Human label: theme (Workload, Pay, Career growth, Manager behaviour, Work location, Leave & time off, Facilities, Team culture, Harassment, Discrimination, Ethics, Positive feedback, plus non-feedback classes)",
  "gold_sub_theme": "Human label: specific sub-theme (lowest level of the generalisation ladder)",
  "gold_domain": "Human label: broad domain (top of the ladder - what senior leadership sees). Idea: generalise complaints as they go higher up",
  "gold_issue_target": "Human label: Policy / Person / Both (idea: split complaints into policy vs person related)",
  "gold_sentiment": "Human label: Negative / Neutral / Positive",
  "gold_severity": "Human label: 1 Low, 2 Medium, 3 High, 4 Critical",
  "gold_severity_label": "Severity as text",
  "gold_escalate": "TRUE if it must go to a human in HR (all Critical cases)",
  "gold_escalation_route": "Where it goes - never the employee's own manager",
  "edge_case_type": "Standard, or the edge case it tests: Re-identifiable after redaction, Self-identifies, Writing-style signature, Sarcasm, Prompt injection, Empty / spam",
  "team_size_at_submission": "Team headcount that month",
  "below_k5_reporting_threshold": "TRUE if team < 5 people - results must not be shown at team level",
  "writer_check_of_rewrite": "Blank on purpose. The app fills it when the writer reviews the anonymised rewrite (Approved / Edited / Meaning changed)",
  "identity_token": "Sealed token, rotated every quarter. Blank for Fully anonymous",
  "receipt_code": "Code the writer keeps to follow up or dispute, without revealing identity",
  "identified_name": "Filled only when the writer chose Identified mode",
  "split": "'LLM test set' = 502 comments (all edge-case types over-sampled) for measuring the LLM against human labels; 'Pool' = the rest",
 },
 "team_monthly_panel": {
  "team_id": "Team ID", "month": "Month (YYYY-MM)", "department": "Department", "location": "Office city",
  "headcount_start": "People in the team at the start of the month",
  "exits_this_month": "People who left this month",
  "hires_this_month": "Backfill hires who joined this month",
  "open_positions": "Unfilled roles at month end",
  "avg_tenure_months": "Average tenure in the team",
  "pct_under_12m_tenure": "% of team with under 12 months tenure",
  "avg_overtime_hrs_week": "Average overtime hours per person per week (from timesheets)",
  "weekend_logins_per_emp": "Average weekend system logins per person",
  "unplanned_leave_days_per_emp": "Average unplanned / sick leave days per person",
  "pulse_score_1to5": "Monthly pulse survey score (1-5)",
  "manager_pulse_score_1to5": "Team's rating of their manager (1-5)",
  "pay_vs_market_pct": "Team's average pay vs market benchmark, % (negative = below market)",
  "last_hike_pct": "Average hike % in the last appraisal cycle",
  "promotions_last_12m": "Promotions in the team in the last 12 months",
  "manager_tenure_months": "How long the current manager has led the team",
  "manager_changed_last_6m": "1 if the manager changed in the last 6 months",
  "internal_transfer_requests": "Requests to move out of the team this month",
  "rto_days_required": "Office days per week required (0 before Apr 2026, 3 after)",
  "avg_commute_km": "Average one-way commute (km)",
  "n_comments": "Feedback comments from this team this month (excludes spam / injection)",
  "n_negative": "Of which negative",
  "severity_sum": "Sum of severity of non-positive comments",
  "n_high_or_critical": "Comments with severity High or Critical",
  "n_escalations": "Comments escalated",
  "n_theme_*": "Count of comments by theme",
  "n_target_*": "Count of comments by Policy / Person / Both",
  "complaint_score_per_10_emp": "Monthly complaint score = severity-weighted complaints per 10 people (idea: monthly complaint score)",
  "negative_share_pct": "% of comments that were negative (blank if no comments)",
  "disputes_raised": "Disputes raised in that month's HR review window (idea: place to argue with HR)",
  "exits_last_3m": "Exits in this month and the two before (the 'historical rate' baseline uses this)",
  "exits_next_3m": "TARGET INPUT - exits in the next 3 months (blank for the last 3 months of data)",
  "attrition_rate_next_3m_pct": "TARGET INPUT - exits_next_3m / headcount_start",
  "target_high_risk_next_3m": "MODEL TARGET - 1 if attrition_rate_next_3m_pct >= 10%",
  "below_k5_reporting_threshold": "TRUE if team < 5 people - do not report at team level; roll up to department",
 },
 "teams": {"team_id": "Team ID", "team_name": "Team name", "department": "Department", "location": "Office city",
           "manager_id": "Manager ID", "manager_name": "Manager name (synthetic) - used to test name redaction",
           "initial_headcount": "Headcount in Nov 2024"},
 "employees_hris": {"employee_id": "Employee ID", "employee_name": "Name (synthetic)", "team_id": "Team",
                    "gender": "Gender - only for bias checks, never a model input", "level": "Grade",
                    "tenure_months_at_start": "Tenure in Nov 2024 (0 for later hires)", "joined_month": "Joining month",
                    "exit_month": "Exit month (blank if active)", "status": "Active / Exited"},
 "listening_sessions": {
  "session_id": "Session ID", "team_id": "Team", "manager_id": "Team's manager", "department": "Department", "location": "City",
  "alarm_month": "Month the alarm fired",
  "alarm_reason": "Why: historical exits >= 10% in the last 3 months (the rule HR can use today) and/or a complaint-score spike (top 5% of 3-month averages). The model's risk score is the third trigger your app adds",
  "exit_rate_last_3m_pct": "Exits in the last 3 months as % of headcount, when the alarm fired",
  "complaint_score_3m": "3-month average complaint score when the alarm fired",
  "session_status": "Held / Not held",
  "held_within_days": "Days from alarm to session",
  "attendees": "People who attended",
  "top_issue_1": "Most common complaint theme in the team's last 3 months (falls back to the worst HR signal if there were no comments)",
  "top_issue_2": "Second most common theme",
  "main_issue_type": "Policy / Person / Both - which kind of issue dominated",
  "actions_agreed": "Actions agreed in the session",
  "actions_closed_on_time": "Of those, closed by the deadline",
  "pulse_before": "Team pulse score when the alarm fired",
  "pulse_avg_next_3m": "Average pulse score in the 3 months after",
  "exits_next_3m": "Exits in the 3 months after the alarm",
 },
 "manager_yearend_review": {
  "review_period": "Appraisal period (April cycle)", "team_id": "Team", "manager_id": "Manager", "manager_name": "Manager name",
  "department": "Department", "avg_headcount": "Average team size in the period",
  "reportable_k5": "FALSE if the team averaged under 5 people - do not use at team level",
  "attrition_annualised_pct": "Exits in the period, annualised, % of headcount",
  "avg_complaint_score": "Average monthly complaint score",
  "complaint_score_change": "Last 3 months minus first 3 months (positive = getting worse)",
  "person_related_complaints_pct": "% of negative comments about a person (Person or Both)",
  "escalations": "Critical comments from the team",
  "alarms": "Listening-session alarms in the period", "sessions_held": "Sessions actually held",
  "actions_agreed": "Actions agreed across sessions", "actions_closed_on_time": "Actions closed on time",
  "actions_closed_pct": "Actions closed on time, %",
  "avg_manager_rating": "Team's average rating of the manager (1-5)",
  "manager_changed_in_period": "TRUE if the manager changed - the period should be shared",
  "team_risk_factor_0_100": "Suggested team risk factor for the manager's review: 30% attrition, 25% complaint score, 15% trend, 15% person-related share, 15% actions not closed. Weights are a starting point for the team to agree. Team level only - never an individual employee's appraisal",
  "review_note": "What to do with the row",
 },
 "hr_review_disputes": {"dispute_id": "Dispute ID", "comment_id": "Comment being contested",
                        "raised_via": "Identity token or receipt code - identity is not revealed",
                        "review_window": "Monthly review window where it was heard",
                        "dispute_reason": "What the employee disagreed with", "outcome": "Reclassified / Upheld / Escalated",
                        "resolved_within_days": "Days to resolve"},
 "identity_vault_SEALED": {"identity_token": "Token shown on comments", "employee_id": "Real employee behind the token",
                           "quarter": "Quarter the token is valid for (rotates quarterly)",
                           "custodian": "Only this role can open the vault (idea: secret authorised identity object beyond HR)",
                           "unseal_allowed_when": "Rule for opening it"},
 "policy_events": {"month": "Month", "event": "What changed", "scope": "Who it affected", "expected_effect": "Expected effect on data"},
 "answer_key_hidden": {"team_id": "Team", "mgr_quality": "Hidden manager quality 0-1 (higher = better)",
                       "workload_base": "Hidden baseline workload", "pay_gap": "Hidden starting pay gap vs market %",
                       "growth_stag": "Hidden career stagnation 0-1", "wfh_dep": "Hidden dependence on WFH 0-1",
                       "commute": "Hidden commute burden 0-1", "toxic_mgr": "TRUE for 5 teams with a problematic manager (harassment cluster)"},
}


def write_table(ws, df, widths=None):
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            ws.cell(row=i, column=j, value=v).font = BODY
    for j, c in enumerate(df.columns, 1):
        w = (widths or {}).get(c)
        if w is None:
            sample = df[c].head(200).map(lambda v: len(str(v)))
            w = min(max(len(c), int(sample.quantile(0.9)) if len(sample) and sample.notna().any() else 8) + 2, 60)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


wb = Workbook()
ws = wb.active
ws.title = "Start here"
cm, pnl = frames["feedback_comments"], frames["team_monthly_panel"]
lines = [
    ("Synthetic dataset - AI-Powered Anonymous Employee Feedback System", TITLE),
    ("Fictional company: Navrang Systems Pvt Ltd, ~2,000 employees, 6 Indian cities, Nov 2024 - Oct 2026. Every name, comment and number is made up.", BODY),
    ("", BODY),
    ("What's inside", BOLD),
    (f"feedback_comments - {len(cm):,} comments with human (gold) labels and correct anonymised versions", BODY),
    (f"team_monthly_panel - {len(pnl):,} team-month rows: HR signals, comment signals, and the attrition target", BODY),
    (f"teams - {len(frames['teams'])} teams (incl. some under 5 people, to test the k>=5 rule)", BODY),
    (f"employees_hris - {len(frames['employees_hris']):,} employee records (incl. leavers and backfills)", BODY),
    (f"listening_sessions - {len(frames['listening_sessions'])} alarms and the sessions HR held after them", BODY),
    (f"manager_yearend_review - {len(frames['manager_yearend_review'])} team x period rows with the team risk factor for the manager's review", BODY),
    (f"hr_review_disputes - {len(frames['hr_review_disputes'])} disputes raised in the monthly HR review window", BODY),
    (f"identity_vault_SEALED - {len(frames['identity_vault_SEALED']):,} token-to-employee mappings (custodian only)", BODY),
    ("policy_events - appraisal cycles, leave policy change, office-return mandate (Apr 2026) and manager reshuffle", BODY),
    ("answer_key_hidden - the hidden drivers used to generate the data. For defending the model in Q&A only - never a model input", BODY),
    ("", BODY),
    ("Which format for what", BOLD),
    ("Tables (this Excel, or the CSVs) = the data itself. The risk model, dashboard and Streamlit app read these.", BODY),
    ("JSON = what goes in and out of the LLM. Give it llm/comments_input_all.jsonl (comment ID + text only), make it reply in the format in llm/llm_output_schema.json, then compare its replies with llm/test_set_gold.jsonl.", BODY),
    ("", BODY),
    ("Where each team idea lives", BOLD),
]
ideas = [
    ("Writer rechecks the LLM's rewrite", "feedback_comments.writer_check_of_rewrite (filled by the app) vs gold_final_text"),
    ("Split into policy / person related", "feedback_comments.gold_issue_target; panel n_target_policy / person / both"),
    ("Risk factor into year-end evaluation", "manager_yearend_review.team_risk_factor_0_100 (team level only, never an individual employee)"),
    ("Monthly complaint score", "team_monthly_panel.complaint_score_per_10_emp"),
    ("Place to argue with HR each month", "hr_review_disputes (raised via token or receipt code)"),
    ("Secret authorised identity beyond HR", "identity_vault_SEALED + feedback_comments.identity_token"),
    ("Choose fully anonymous or not", "feedback_comments.anonymity_mode"),
    ("Generalise complaints as they go higher up", "gold_sub_theme -> gold_theme -> gold_domain"),
    ("Date anonymous as a window", "feedback_comments.submission_window (half-month, no exact date)"),
    ("Retraining every few months", "Office-return mandate in Apr 2026 changes what drives attrition; train before vs after"),
    ("Two layers of privacy", "raw_text -> gold_layer1_text (names out) -> gold_final_text (indirect clues out); reid_risk_after_layer1"),
    ("Historical-risk threshold alarm -> listening sessions", "listening_sessions (alarm, session, actions, what happened after); team_monthly_panel.exits_last_3m and target_high_risk_next_3m for the model"),
    ("LLM anonymises and classifies (JSON)", "llm/ folder: comments_input_all.jsonl, test_set_input.jsonl, test_set_gold.jsonl, llm_output_schema.json"),
    ("Leak check after anonymising", "llm/redaction_lists.json (names, IDs, project codes, clients, offices)"),
    ("Escalation (harassment, discrimination, ethics)", "feedback_comments.gold_escalate + gold_escalation_route"),
    ("Minimum group of 5", "below_k5_reporting_threshold (comments, panel); reportable_k5 (year-end)"),
    ("Explain reasons behind high risk", "HR signal columns (overtime, pay vs market, manager score, promotions, office days) + theme shares"),
    ("Demo cases (routine / anonymise / escalate / edge)", "filter feedback_comments by edge_case_type, contains_direct_pii, gold_escalate"),
]
r = 1
for text, f in lines:
    ws.cell(row=r, column=1, value=text).font = f
    r += 1
ws.cell(row=r, column=1, value="Idea").font = HDR_FONT
ws.cell(row=r, column=1).fill = HDR_FILL
ws.cell(row=r, column=2, value="Where to find it").font = HDR_FONT
ws.cell(row=r, column=2).fill = HDR_FILL
r += 1
for a, b in ideas:
    ws.cell(row=r, column=1, value=a).font = BODY
    ws.cell(row=r, column=2, value=b).font = BODY
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
r += 1
ws.cell(row=r, column=1, value="Key numbers").font = BOLD
r += 1
months = pnl.groupby("month").headcount_start.sum().mean()
facts = [
    ("Annual attrition (approx.)", f"{round(100 * pnl.exits_this_month.sum() / 2 / months)}%"),
    ("Team-months labelled high risk (>=10% exits in next 3 months)", f"{round(100 * pnl.target_high_risk_next_3m.mean())}%"),
    ("Comments with direct identifiers", f"{int(cm.contains_direct_pii.sum()):,}"),
    ("Comments re-identifiable after names removed (Medium/High)", f"{int((cm.reid_risk_after_layer1 != 'Low').sum()):,}"),
    ("Comments that must be escalated", f"{int(cm.gold_escalate.sum())}"),
    ("Prompt-injection attempts planted", f"{int((cm.edge_case_type == 'Prompt injection').sum())}"),
    ("Hinglish comments", f"{int((cm.language == 'Hinglish').sum())}"),
]
for a, b in facts:
    ws.cell(row=r, column=1, value=a).font = BODY
    ws.cell(row=r, column=2, value=b).font = BODY
    r += 1
r += 1
ws.cell(row=r, column=1, value="Do not put employee_id_hidden, identity_vault_SEALED or answer_key_hidden into the LLM prompt or the model.").font = Font(name=ARIAL, size=10, bold=True, color="9C0006")
ws.cell(row=r, column=1).fill = WARN_FILL
ws.column_dimensions["A"].width = 58
ws.column_dimensions["B"].width = 95

# dictionary
wd = wb.create_sheet("Data dictionary")
rows = [(s, c, d) for s, cols in DICT.items() for c, d in cols.items()]
write_table(wd, pd.DataFrame(rows, columns=["sheet", "column", "meaning"]), {"sheet": 24, "column": 32, "meaning": 120})
for row in wd.iter_rows(min_row=2):
    row[2].alignment = Alignment(wrap_text=True, vertical="top")

for name, _ in sheets:
    wsx = wb.create_sheet(name)
    widths = {"raw_text": 70, "gold_layer1_text": 70, "gold_final_text": 70, "unseal_allowed_when": 60,
              "expected_effect": 70, "event": 45, "alarm_reason": 40, "review_note": 45}
    write_table(wsx, frames[name], widths)
    if name in ("identity_vault_SEALED", "answer_key_hidden"):
        wsx.sheet_properties.tabColor = "C00000"

wb.save(OUT)
print("saved", OUT)
